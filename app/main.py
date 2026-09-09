from datetime import date, datetime
from io import BytesIO
import os
import re
from statistics import median
import sys
from pathlib import Path

# Allow running both as module: `python -m app.main`
# and as file: `python app/main.py`
if __package__ is None or __package__ == "":
    sys.path.append(os.path.dirname(os.path.dirname(__file__)))

from flask import Flask, abort, redirect, render_template, request, send_file, url_for
from sqlalchemy import case, func, or_
from sqlalchemy.orm import joinedload
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.database import Base, engine, get_session
from app.models import Pallet, Shoe, ShoePhoto, VintedAccount


app = Flask(__name__, static_folder="static", template_folder="templates")

Base.metadata.create_all(bind=engine)


def ensure_schema_compatibility() -> None:
    with engine.begin() as conn:
        columns = [row[1] for row in conn.exec_driver_sql("PRAGMA table_info(shoes)").fetchall()]
        if "sale_price" not in columns:
            conn.exec_driver_sql("ALTER TABLE shoes ADD COLUMN sale_price FLOAT")
        if "sale_date" not in columns:
            conn.exec_driver_sql("ALTER TABLE shoes ADD COLUMN sale_date DATE")
        if "vinted_account_id" not in columns:
            conn.exec_driver_sql("ALTER TABLE shoes ADD COLUMN vinted_account_id INTEGER")


ensure_schema_compatibility()

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
PHOTOS_ROOT = Path.cwd() / "PALETY"
PHOTOS_ROOT.mkdir(parents=True, exist_ok=True)


def expected_shoe_dir(pallet_code: str, shoe_internal_id: str) -> Path:
    return PHOTOS_ROOT / pallet_code / shoe_internal_id


def next_internal_id_for_pallet(pallet_code: str, internal_ids: list[str]) -> str:
    raw_prefix = (pallet_code.split("_")[-1] if "_" in pallet_code else pallet_code).strip()
    prefix = raw_prefix.upper() if raw_prefix else "X"
    pattern = re.compile(rf"^{re.escape(prefix)}(\d+)$", re.IGNORECASE)

    max_num = 0
    for item in internal_ids:
        match = pattern.match((item or "").strip())
        if match:
            max_num = max(max_num, int(match.group(1)))

    return f"{prefix}{max_num + 1:03d}"


def to_rel_path(path_obj: Path) -> str:
    return str(path_obj.relative_to(Path.cwd())).replace("\\", "/")


def is_supported_image(path_obj: Path) -> bool:
    return path_obj.is_file() and path_obj.suffix.lower() in IMAGE_EXTENSIONS


def sync_shoe_photos(session, pallet: Pallet, shoe: Shoe) -> None:
    folder = expected_shoe_dir(pallet.code, shoe.internal_id)
    if not folder.exists() or not folder.is_dir():
        legacy_folder = Path.cwd() / pallet.code / shoe.internal_id
        if legacy_folder.exists() and legacy_folder.is_dir():
            folder = legacy_folder
        else:
            return

    disk_photos = sorted([p for p in folder.iterdir() if is_supported_image(p)], key=lambda x: x.name.lower())
    disk_rel_paths = [to_rel_path(p) for p in disk_photos]

    existing = session.query(ShoePhoto).filter(ShoePhoto.shoe_id == shoe.id).all()
    existing_map = {photo.file_path: photo for photo in existing}

    for index, rel_path in enumerate(disk_rel_paths):
        if rel_path in existing_map:
            existing_map[rel_path].sort_order = index
        else:
            session.add(ShoePhoto(shoe_id=shoe.id, file_path=rel_path, sort_order=index))

    for old_path, photo in existing_map.items():
        if old_path not in set(disk_rel_paths):
            session.delete(photo)


def safe_abs_from_rel(rel_path: str) -> Path | None:
    base = Path.cwd().resolve()
    target = (base / rel_path).resolve()
    if str(target).startswith(str(base)):
        return target
    return None


def month_date_range(today: date | None = None) -> tuple[date, date]:
    current = today or date.today()
    start = current.replace(day=1)
    if current.month == 12:
        next_month_start = date(current.year + 1, 1, 1)
    else:
        next_month_start = date(current.year, current.month + 1, 1)
    end = next_month_start.fromordinal(next_month_start.toordinal() - 1)
    return start, end


def local_today() -> date:
    return datetime.now().astimezone().date()


def previous_month_date_range(today: date | None = None) -> tuple[date, date]:
    current = today or local_today()
    current_month_start = current.replace(day=1)
    previous_month_end = current_month_start.fromordinal(current_month_start.toordinal() - 1)
    return previous_month_end.replace(day=1), previous_month_end


def parse_percent_rate(raw_value: str, default: float = 3.0) -> float:
    if not raw_value:
        return default
    cleaned = raw_value.strip().replace(",", ".")
    try:
        value = float(cleaned)
    except ValueError:
        return default
    if value < 0:
        return default
    return value


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/media/<path:rel_path>")
def media_file(rel_path: str):
    abs_path = safe_abs_from_rel(rel_path)
    if not abs_path or not abs_path.exists() or not abs_path.is_file():
        abort(404)
    return send_file(abs_path)


@app.get("/")
def home():
    month_start, month_end = month_date_range()
    start_raw = request.args.get("report_start", "").strip()
    end_raw = request.args.get("report_end", "").strip()
    rate_raw = request.args.get("tax_rate", "").strip()
    shoe_q = request.args.get("shoe_q", "").strip()

    try:
        report_start = datetime.strptime(start_raw, "%Y-%m-%d").date() if start_raw else month_start
    except ValueError:
        report_start = month_start

    try:
        report_end = datetime.strptime(end_raw, "%Y-%m-%d").date() if end_raw else month_end
    except ValueError:
        report_end = month_end

    if report_end < report_start:
        report_end = report_start

    tax_rate = parse_percent_rate(rate_raw, default=3.0)

    with get_session() as session:
        pallets = session.query(Pallet).order_by(Pallet.id.desc()).all()

        sales_by_pallet = {
            pallet_id: {
                "total_count": total_count,
                "sold_count": sold_count,
                "sold_value": float(sold_value or 0.0),
            }
            for pallet_id, total_count, sold_count, sold_value in (
                session.query(
                    Shoe.pallet_id,
                    func.count(Shoe.id).label("total_count"),
                    func.sum(case((Shoe.status == "sold", 1), else_=0)).label("sold_count"),
                    func.sum(case((Shoe.status == "sold", Shoe.sale_price), else_=0.0)).label("sold_value"),
                )
                .group_by(Shoe.pallet_id)
                .all()
            )
        }

        pallet_returns = {}
        for pallet in pallets:
            sales = sales_by_pallet.get(pallet.id, {"total_count": 0, "sold_count": 0, "sold_value": 0.0})
            purchase_cost = float(pallet.purchase_gross or 0.0)
            sold_value = sales["sold_value"]
            return_percent = (sold_value / purchase_cost * 100.0) if purchase_cost > 0 else 0.0
            pallet_returns[pallet.id] = {
                "total_count": sales["total_count"],
                "sold_count": sales["sold_count"],
                "sold_value": sold_value,
                "return_percent": return_percent,
                "bar_percent": min(return_percent, 100.0),
            }

        all_shoes = (
            session.query(Shoe)
            .options(joinedload(Shoe.pallet), joinedload(Shoe.vinted_account))
            .order_by(Shoe.id.desc())
            .all()
        )
        sold_shoes = [shoe for shoe in all_shoes if shoe.status == "sold"]
        available_shoes = [shoe for shoe in all_shoes if shoe.status != "sold"]
        sold_prices = [float(shoe.sale_price) for shoe in sold_shoes if shoe.sale_price is not None]

        total_purchase_cost = sum(float(pallet.purchase_gross or 0.0) for pallet in pallets)
        total_sales_value = sum(sold_prices)
        financial_surplus = total_sales_value - total_purchase_cost
        global_return_percent = (
            total_sales_value / total_purchase_cost * 100.0 if total_purchase_cost > 0 else 0.0
        )
        sold_share_percent = len(sold_shoes) / len(all_shoes) * 100.0 if all_shoes else 0.0

        inventory_summary = {
            "total_purchase_cost": total_purchase_cost,
            "total_sales_value": total_sales_value,
            "financial_surplus": financial_surplus,
            "global_return_percent": global_return_percent,
            "total_shoes": len(all_shoes),
            "sold_shoes": len(sold_shoes),
            "available_shoes": len(available_shoes),
            "sold_share_percent": sold_share_percent,
        }

        price_statistics = {
            "average": sum(sold_prices) / len(sold_prices) if sold_prices else 0.0,
            "median": float(median(sold_prices)) if sold_prices else 0.0,
            "minimum": min(sold_prices) if sold_prices else 0.0,
            "maximum": max(sold_prices) if sold_prices else 0.0,
        }

        top_sold_shoes = sorted(
            (shoe for shoe in sold_shoes if shoe.sale_price is not None),
            key=lambda shoe: (float(shoe.sale_price or 0.0), shoe.sale_date or date.min),
            reverse=True,
        )[:10]

        today = local_today()
        current_month_start, current_month_end = month_date_range(today)
        previous_month_start, previous_month_end = previous_month_date_range(today)

        def sales_for_period(start: date, end: date) -> dict:
            period_shoes = [
                shoe for shoe in sold_shoes if shoe.sale_date and start <= shoe.sale_date <= end
            ]
            return {
                "value": sum(float(shoe.sale_price or 0.0) for shoe in period_shoes),
                "count": len(period_shoes),
            }

        current_month_sales = sales_for_period(current_month_start, current_month_end)
        previous_month_sales = sales_for_period(previous_month_start, previous_month_end)
        if previous_month_sales["value"] > 0:
            monthly_change_percent = (
                (current_month_sales["value"] - previous_month_sales["value"])
                / previous_month_sales["value"]
                * 100.0
            )
        else:
            monthly_change_percent = None

        polish_months = (
            "styczeń", "luty", "marzec", "kwiecień", "maj", "czerwiec",
            "lipiec", "sierpień", "wrzesień", "październik", "listopad", "grudzień",
        )
        monthly_statistics = {
            "current": current_month_sales,
            "previous": previous_month_sales,
            "current_label": f"{polish_months[current_month_start.month - 1]} {current_month_start.year}",
            "previous_label": f"{polish_months[previous_month_start.month - 1]} {previous_month_start.year}",
            "change_percent": monthly_change_percent,
        }

        oldest_available_shoes = []
        for shoe in sorted(
            available_shoes,
            key=lambda item: item.created_at or datetime.max,
        )[:10]:
            created_date = shoe.created_at.date() if shoe.created_at else today
            oldest_available_shoes.append({
                "shoe": shoe,
                "days_in_inventory": max((today - created_date).days, 0),
            })

        attention_pallets = []
        for pallet in pallets:
            return_data = pallet_returns[pallet.id]
            age_days = max((today - pallet.delivery_date).days, 0)
            if return_data["return_percent"] < 50.0 and age_days >= 30 and return_data["total_count"] > 0:
                attention_pallets.append({
                    "pallet": pallet,
                    "age_days": age_days,
                    "return_percent": return_data["return_percent"],
                    "sold_count": return_data["sold_count"],
                    "total_count": return_data["total_count"],
                    "missing_to_return": max(
                        float(pallet.purchase_gross or 0.0) - return_data["sold_value"], 0.0
                    ),
                })
        attention_pallets.sort(key=lambda item: (item["return_percent"], -item["age_days"]))
        attention_pallets = attention_pallets[:8]

        dashboard_shoe_ids = [shoe.id for shoe in top_sold_shoes]
        dashboard_photos = (
            session.query(ShoePhoto)
            .filter(ShoePhoto.shoe_id.in_(dashboard_shoe_ids))
            .order_by(ShoePhoto.sort_order.asc(), ShoePhoto.id.asc())
            .all()
            if dashboard_shoe_ids
            else []
        )
        dashboard_first_photo_by_shoe = {}
        for photo in dashboard_photos:
            dashboard_first_photo_by_shoe.setdefault(photo.shoe_id, photo)

        global_shoes = []
        first_photo_by_shoe = {}
        if shoe_q:
            global_shoes = (
                session.query(Shoe)
                .options(joinedload(Shoe.pallet))
                .join(Pallet)
                .filter(
                    or_(
                        Shoe.internal_id.ilike(f"%{shoe_q}%"),
                        Shoe.name.ilike(f"%{shoe_q}%"),
                        Pallet.code.ilike(f"%{shoe_q}%"),
                    )
                )
                .order_by(Shoe.id.desc())
                .limit(100)
                .all()
            )

            shoe_ids = [s.id for s in global_shoes]
            photos = (
                session.query(ShoePhoto)
                .filter(ShoePhoto.shoe_id.in_(shoe_ids))
                .order_by(ShoePhoto.sort_order.asc(), ShoePhoto.id.asc())
                .all()
                if shoe_ids
                else []
            )

            photos_by_shoe = {}
            for photo in photos:
                photos_by_shoe.setdefault(photo.shoe_id, []).append(photo)

            first_photo_by_shoe = {
                shoe_id: shoe_photos[0] for shoe_id, shoe_photos in photos_by_shoe.items() if shoe_photos
            }

        sold_in_range = (
            session.query(Shoe)
            .filter(
                Shoe.status == "sold",
                Shoe.sale_date.isnot(None),
                Shoe.sale_date >= report_start,
                Shoe.sale_date <= report_end,
            )
            .all()
        )

        report_total_sales = sum((s.sale_price or 0.0) for s in sold_in_range)

    return render_template(
        "home.html",
        pallets=pallets,
        report_start=report_start.isoformat(),
        report_end=report_end.isoformat(),
        tax_rate=tax_rate,
        report_total_sales=report_total_sales,
        shoe_q=shoe_q,
        global_shoes=global_shoes,
        first_photo_by_shoe=first_photo_by_shoe,
        pallet_returns=pallet_returns,
        inventory_summary=inventory_summary,
        price_statistics=price_statistics,
        top_sold_shoes=top_sold_shoes,
        dashboard_first_photo_by_shoe=dashboard_first_photo_by_shoe,
        monthly_statistics=monthly_statistics,
        oldest_available_shoes=oldest_available_shoes,
        attention_pallets=attention_pallets,
    )


@app.post("/")
def home_post():
    return redirect(url_for("home"), code=303)


@app.get("/reports/sold-shoes.xlsx")
def export_sold_shoes_report():
    month_start, month_end = month_date_range()
    start_raw = request.args.get("report_start", "").strip()
    end_raw = request.args.get("report_end", "").strip()
    rate_raw = request.args.get("tax_rate", "").strip()

    try:
        report_start = datetime.strptime(start_raw, "%Y-%m-%d").date() if start_raw else month_start
    except ValueError:
        report_start = month_start

    try:
        report_end = datetime.strptime(end_raw, "%Y-%m-%d").date() if end_raw else month_end
    except ValueError:
        report_end = month_end

    if report_end < report_start:
        report_end = report_start

    tax_rate = parse_percent_rate(rate_raw, default=3.0)
    tax_multiplier = tax_rate / 100.0

    with get_session() as session:
        sold_shoes = (
            session.query(Shoe)
            .filter(
                Shoe.status == "sold",
                Shoe.sale_date.isnot(None),
                Shoe.sale_date >= report_start,
                Shoe.sale_date <= report_end,
            )
            .order_by(Shoe.sale_date.asc(), Shoe.id.asc())
            .all()
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sprzedane buty"

    headers = [
        "ID rekordu",
        "Data sprzedaży",
        "ID buta",
        "Nazwa",
        "Cena sprzedaży",
        "Stawka ryczałtu",
        "Podatek",
    ]
    sheet.append(headers)

    header_fill = PatternFill(fill_type="solid", start_color="1F2937", end_color="1F2937")
    header_font = Font(color="FFFFFF", bold=True)

    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for shoe in sold_shoes:
        sale_price = float(shoe.sale_price or 0.0)
        tax_value = sale_price * tax_multiplier
        sheet.append(
            [
                shoe.id,
                shoe.sale_date.isoformat() if shoe.sale_date else "",
                shoe.internal_id,
                shoe.name,
                sale_price,
                tax_rate / 100.0,
                tax_value,
            ]
        )

    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=7, max_row=sheet.max_row):
        row[0].number_format = '#,##0.00 "zł"'
        row[1].number_format = "0.00%"
        row[2].number_format = '#,##0.00 "zł"'

    data_last_row = sheet.max_row
    if data_last_row >= 2:
        summary_row = data_last_row + 2
        sheet.cell(row=summary_row, column=4, value="SUMA SPRZEDAŻY").font = Font(bold=True)
        total_cell = sheet.cell(row=summary_row, column=5, value=f"=SUM(E2:E{data_last_row})")
        total_cell.number_format = '#,##0.00 "zł"'
        total_cell.font = Font(bold=True)

        sheet.cell(row=summary_row + 1, column=4, value="SUMA PODATKU").font = Font(bold=True)
        tax_total_cell = sheet.cell(row=summary_row + 1, column=5, value=f"=SUM(G2:G{data_last_row})")
        tax_total_cell.number_format = '#,##0.00 "zł"'
        tax_total_cell.font = Font(bold=True)

    column_widths = {
        1: 12,
        2: 16,
        3: 14,
        4: 36,
        5: 18,
        6: 18,
        7: 16,
    }
    for col_idx, width in column_widths.items():
        sheet.column_dimensions[chr(64 + col_idx)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"raport_sprzedanych_butow_{report_start.isoformat()}_{report_end.isoformat()}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/db-inspector")
def db_inspector():
    with get_session() as session:
        pallets = session.query(Pallet).order_by(Pallet.id.asc()).all()
        shoes = session.query(Shoe).options(joinedload(Shoe.vinted_account)).order_by(Shoe.id.asc()).all()
        photos = session.query(ShoePhoto).order_by(ShoePhoto.id.asc()).all()
        vinted_accounts = session.query(VintedAccount).order_by(VintedAccount.id.asc()).all()

        stats = {
            "pallets_count": len(pallets),
            "shoes_count": len(shoes),
            "photos_count": len(photos),
            "vinted_accounts_count": len(vinted_accounts),
            "sold_shoes_count": sum(1 for s in shoes if s.status == "sold"),
            "sold_total": sum((s.sale_price or 0.0) for s in shoes if s.status == "sold"),
        }

    return render_template(
        "db_inspector.html",
        pallets=pallets,
        shoes=shoes,
        photos=photos,
        vinted_accounts=vinted_accounts,
        stats=stats,
    )


@app.get("/vinted-accounts")
def vinted_accounts():
    with get_session() as session:
        accounts = session.query(VintedAccount).order_by(VintedAccount.name.asc()).all()

    return render_template("vinted_accounts.html", accounts=accounts)


@app.post("/vinted-accounts")
def create_vinted_account():
    name = request.form.get("name", "").strip()
    is_banned = request.form.get("is_banned") == "on"

    if not name:
        return redirect(url_for("vinted_accounts"))

    with get_session() as session:
        exists = session.query(VintedAccount).filter(VintedAccount.name == name).first()
        if not exists:
            session.add(VintedAccount(name=name, is_banned=is_banned))
            session.commit()

    return redirect(url_for("vinted_accounts"))


@app.post("/vinted-accounts/<int:account_id>")
def update_vinted_account(account_id: int):
    name = request.form.get("name", "").strip()
    is_banned = request.form.get("is_banned") == "on"

    if not name:
        return redirect(url_for("vinted_accounts"))

    with get_session() as session:
        account = session.query(VintedAccount).filter(VintedAccount.id == account_id).first()
        if not account:
            return redirect(url_for("vinted_accounts"))

        duplicate = (
            session.query(VintedAccount)
            .filter(VintedAccount.name == name, VintedAccount.id != account_id)
            .first()
        )
        if duplicate:
            return redirect(url_for("vinted_accounts"))

        account.name = name
        account.is_banned = is_banned
        session.commit()

    return redirect(url_for("vinted_accounts"))


@app.get("/pallets/<int:pallet_id>")
def pallet_detail(pallet_id: int):
    query_text = request.args.get("q", "").strip()
    status_filter = request.args.get("status", "all").strip().lower()

    with get_session() as session:
        pallet = session.query(Pallet).filter(Pallet.id == pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        vinted_accounts = session.query(VintedAccount).order_by(VintedAccount.is_banned.asc(), VintedAccount.name.asc()).all()

        shoes_query = session.query(Shoe).options(joinedload(Shoe.vinted_account)).filter(Shoe.pallet_id == pallet_id)
        if status_filter in {"available", "sold"}:
            shoes_query = shoes_query.filter(Shoe.status == status_filter)
        if query_text:
            shoes_query = shoes_query.filter(
                or_(
                    Shoe.internal_id.ilike(f"%{query_text}%"),
                    Shoe.name.ilike(f"%{query_text}%"),
                )
            )

        shoes = shoes_query.order_by(Shoe.id.desc()).all()

        all_shoes_in_pallet = session.query(Shoe).filter(Shoe.pallet_id == pallet_id).all()
        default_internal_id = next_internal_id_for_pallet(
            pallet.code,
            [s.internal_id for s in all_shoes_in_pallet],
        )
        total_shoes_count = len(all_shoes_in_pallet)
        sold_shoes_count = sum(1 for item in all_shoes_in_pallet if item.status == "sold")
        sold_value = sum((item.sale_price or 0.0) for item in all_shoes_in_pallet if item.status == "sold")
        sold_value_label = f"{sold_value:.2f} zł"

        shoe_ids = [s.id for s in shoes]
        photos = (
            session.query(ShoePhoto)
            .filter(ShoePhoto.shoe_id.in_(shoe_ids))
            .order_by(ShoePhoto.sort_order.asc(), ShoePhoto.id.asc())
            .all()
            if shoe_ids
            else []
        )

        photos_by_shoe = {}
        for photo in photos:
            photos_by_shoe.setdefault(photo.shoe_id, []).append(photo)

        first_photo_by_shoe = {
            shoe_id: shoe_photos[0] for shoe_id, shoe_photos in photos_by_shoe.items() if shoe_photos
        }

        expected_dirs = {s.id: str(expected_shoe_dir(pallet.code, s.internal_id)) for s in shoes}

    return render_template(
        "pallet_detail.html",
        pallet=pallet,
        shoes=shoes,
        q=query_text,
        status_filter=status_filter,
        first_photo_by_shoe=first_photo_by_shoe,
        default_internal_id=default_internal_id,
        total_shoes_count=total_shoes_count,
        sold_shoes_count=sold_shoes_count,
        sold_value_label=sold_value_label,
        vinted_accounts=vinted_accounts,
    )


@app.get("/shoes/<int:shoe_id>")
def shoe_detail(shoe_id: int):
    with get_session() as session:
        shoe = session.query(Shoe).options(joinedload(Shoe.vinted_account)).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        pallet = session.query(Pallet).filter(Pallet.id == shoe.pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        photos = (
            session.query(ShoePhoto)
            .filter(ShoePhoto.shoe_id == shoe.id)
            .order_by(ShoePhoto.sort_order.asc(), ShoePhoto.id.asc())
            .all()
        )

        expected_dir = str(expected_shoe_dir(pallet.code, shoe.internal_id))
        vinted_accounts = session.query(VintedAccount).order_by(VintedAccount.is_banned.asc(), VintedAccount.name.asc()).all()

    return render_template(
        "shoe_detail.html",
        shoe=shoe,
        pallet=pallet,
        photos=photos,
        expected_dir=expected_dir,
        default_sale_date=local_today().isoformat(),
        vinted_accounts=vinted_accounts,
    )


@app.post("/pallets")
def create_pallet():
    code = request.form["code"].strip()
    delivery_date_raw = request.form["delivery_date"]
    purchase_gross = float(request.form["purchase_gross"])
    notes = request.form.get("notes", "").strip() or None

    delivery_date = datetime.strptime(delivery_date_raw, "%Y-%m-%d").date()

    with get_session() as session:
        pallet = Pallet(
            code=code,
            delivery_date=delivery_date,
            purchase_gross=purchase_gross,
            notes=notes,
        )
        session.add(pallet)
        session.commit()

    return redirect(url_for("home"))


@app.post("/pallets/<int:pallet_id>/delete")
def delete_pallet(pallet_id: int):
    with get_session() as session:
        pallet = session.query(Pallet).filter(Pallet.id == pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        session.delete(pallet)
        session.commit()

    return redirect(url_for("home"))


@app.post("/pallets/<int:pallet_id>/price")
def update_pallet_price(pallet_id: int):
    price_raw = request.form.get("purchase_gross", "").strip().replace(",", ".")
    try:
        price = float(price_raw)
    except ValueError:
        return redirect(url_for("pallet_detail", pallet_id=pallet_id))

    if price < 0:
        return redirect(url_for("pallet_detail", pallet_id=pallet_id))

    with get_session() as session:
        pallet = session.query(Pallet).filter(Pallet.id == pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        pallet.purchase_gross = price
        session.commit()

    return redirect(url_for("pallet_detail", pallet_id=pallet_id))


@app.post("/pallets/<int:pallet_id>/notes")
def update_pallet_notes(pallet_id: int):
    notes = request.form.get("notes", "").strip() or None

    with get_session() as session:
        pallet = session.query(Pallet).filter(Pallet.id == pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        pallet.notes = notes
        session.commit()

    return redirect(url_for("pallet_detail", pallet_id=pallet_id))


@app.post("/pallets/<int:pallet_id>/shoes")
def create_shoe(pallet_id: int):
    internal_id = request.form["internal_id"].strip()
    name = request.form["name"].strip()
    description = request.form.get("description", "").strip() or None
    vinted_account_raw = request.form.get("vinted_account_id", "").strip()
    vinted_account_id = int(vinted_account_raw) if vinted_account_raw.isdigit() else None

    if not internal_id or not name:
        return redirect(url_for("pallet_detail", pallet_id=pallet_id))

    with get_session() as session:
        exists = session.query(Shoe).filter(Shoe.internal_id == internal_id).first()
        if exists:
            return redirect(url_for("pallet_detail", pallet_id=pallet_id, q=internal_id))

        if vinted_account_id is not None:
            account = session.query(VintedAccount).filter(VintedAccount.id == vinted_account_id).first()
            if not account:
                vinted_account_id = None

        shoe = Shoe(
            pallet_id=pallet_id,
            internal_id=internal_id,
            name=name,
            description=description,
            vinted_account_id=vinted_account_id,
            status="available",
        )
        session.add(shoe)
        session.commit()

    return redirect(url_for("pallet_detail", pallet_id=pallet_id))


@app.post("/shoes/<int:shoe_id>/details")
def update_shoe_details(shoe_id: int):
    internal_id = request.form.get("internal_id", "").strip()
    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip() or None
    vinted_account_raw = request.form.get("vinted_account_id", "").strip()
    vinted_account_id = int(vinted_account_raw) if vinted_account_raw.isdigit() else None

    if not internal_id or not name:
        return redirect(url_for("shoe_detail", shoe_id=shoe_id))

    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        duplicate = session.query(Shoe).filter(Shoe.internal_id == internal_id, Shoe.id != shoe_id).first()
        if duplicate:
            return redirect(url_for("shoe_detail", shoe_id=shoe_id))

        if vinted_account_id is not None:
            account = session.query(VintedAccount).filter(VintedAccount.id == vinted_account_id).first()
            if not account:
                vinted_account_id = None

        shoe.internal_id = internal_id
        shoe.name = name
        shoe.description = description
        shoe.vinted_account_id = vinted_account_id
        session.commit()

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


@app.post("/pallets/<int:pallet_id>/photos/sync-all")
def sync_all_photos_in_pallet(pallet_id: int):
    with get_session() as session:
        pallet = session.query(Pallet).filter(Pallet.id == pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        shoes = session.query(Shoe).filter(Shoe.pallet_id == pallet_id).all()
        for shoe in shoes:
            sync_shoe_photos(session, pallet, shoe)

        session.commit()

    return redirect(url_for("pallet_detail", pallet_id=pallet_id))


@app.post("/shoes/<int:shoe_id>/toggle-status")
def toggle_shoe_status(shoe_id: int):
    sale_price_raw = request.form.get("sale_price", "").strip()
    sale_date_raw = request.form.get("sale_date", "").strip()
    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        if shoe.status == "available":
            if not sale_price_raw:
                return redirect(url_for("shoe_detail", shoe_id=shoe_id))
            try:
                sale_price = float(sale_price_raw.replace(",", "."))
            except ValueError:
                return redirect(url_for("shoe_detail", shoe_id=shoe_id))

            if sale_price < 0:
                return redirect(url_for("shoe_detail", shoe_id=shoe_id))

            if not sale_date_raw:
                return redirect(url_for("shoe_detail", shoe_id=shoe_id))

            try:
                sale_date = datetime.strptime(sale_date_raw, "%Y-%m-%d").date()
            except ValueError:
                return redirect(url_for("shoe_detail", shoe_id=shoe_id))

            shoe.status = "sold"
            shoe.sale_price = sale_price
            shoe.sale_date = sale_date
        else:
            shoe.status = "available"
            shoe.sale_price = None
            shoe.sale_date = None

        pallet_id = shoe.pallet_id
        session.commit()

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


@app.post("/shoes/<int:shoe_id>/sale")
def update_shoe_sale_details(shoe_id: int):
    sale_price_raw = request.form.get("sale_price", "").strip().replace(",", ".")
    sale_date_raw = request.form.get("sale_date", "").strip()

    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        if shoe.status != "sold":
            return redirect(url_for("shoe_detail", shoe_id=shoe_id))

        if not sale_price_raw or not sale_date_raw:
            return redirect(url_for("shoe_detail", shoe_id=shoe_id))

        try:
            sale_price = float(sale_price_raw)
            sale_date = datetime.strptime(sale_date_raw, "%Y-%m-%d").date()
        except ValueError:
            return redirect(url_for("shoe_detail", shoe_id=shoe_id))

        if sale_price < 0:
            return redirect(url_for("shoe_detail", shoe_id=shoe_id))

        shoe.sale_price = sale_price
        shoe.sale_date = sale_date
        session.commit()

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


@app.post("/shoes/<int:shoe_id>/delete")
def delete_shoe(shoe_id: int):
    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        pallet_id = shoe.pallet_id
        session.delete(shoe)
        session.commit()

    return redirect(url_for("pallet_detail", pallet_id=pallet_id))


@app.post("/shoes/<int:shoe_id>/photos/sync")
def sync_photos_for_shoe(shoe_id: int):
    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        pallet = session.query(Pallet).filter(Pallet.id == shoe.pallet_id).first()
        if not pallet:
            return redirect(url_for("home"))

        sync_shoe_photos(session, pallet, shoe)
        session.commit()
        pallet_id = pallet.id

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


@app.post("/shoes/<int:shoe_id>/photos")
def add_photo_path(shoe_id: int):
    file_path = request.form.get("file_path", "").strip().replace("\\", "/")
    with get_session() as session:
        shoe = session.query(Shoe).filter(Shoe.id == shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        if file_path:
            exists = (
                session.query(ShoePhoto)
                .filter(ShoePhoto.shoe_id == shoe.id, ShoePhoto.file_path == file_path)
                .first()
            )
            if not exists:
                max_order = (
                    session.query(ShoePhoto.sort_order)
                    .filter(ShoePhoto.shoe_id == shoe.id)
                    .order_by(ShoePhoto.sort_order.desc())
                    .first()
                )
                next_order = (max_order[0] + 1) if max_order else 0
                session.add(ShoePhoto(shoe_id=shoe.id, file_path=file_path, sort_order=next_order))
                session.commit()

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


@app.post("/shoe-photos/<int:photo_id>/delete")
def delete_photo_path(photo_id: int):
    with get_session() as session:
        photo = session.query(ShoePhoto).filter(ShoePhoto.id == photo_id).first()
        if not photo:
            return redirect(url_for("home"))

        shoe = session.query(Shoe).filter(Shoe.id == photo.shoe_id).first()
        if not shoe:
            return redirect(url_for("home"))

        shoe_id = shoe.id
        session.delete(photo)
        session.commit()

    return redirect(url_for("shoe_detail", shoe_id=shoe_id))


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)

